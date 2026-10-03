# -*- coding: utf-8 -*-
import io
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_attendance_report(start_date_str: str, end_date_str: str, all_students: list, records: list) -> io.BytesIO:
    start_dt = datetime.strptime(start_date_str, "%Y-%m-%d").date()
    end_dt = datetime.strptime(end_date_str, "%Y-%m-%d").date()

    # Формируем список всех дат в выбранном диапазоне
    days_count = (end_dt - start_dt).days + 1
    period_dates = [(start_dt + timedelta(days=i)).isoformat() for i in range(days_count)]

    # Индексируем записи из БД: (student_name, date) -> (status, time)
    attendance_map = {}
    for student, rec_date, rec_time, status in records:
        attendance_map[(student, rec_date)] = {"status": status, "time": rec_time}

    wb = Workbook()

    # ---------------- Стили оформления ----------------
    font_title = Font(name="Calibri", size=14, bold=True, color="1F2937")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="4B5563")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="111827")
    font_regular = Font(name="Calibri", size=11, color="111827")

    font_present = Font(name="Calibri", size=11, bold=True, color="065F46")
    font_absent = Font(name="Calibri", size=11, color="9CA3AF")

    fill_header = PatternFill(start_color="374151", end_color="374151", fill_type="solid")
    fill_summary_header = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    fill_present = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_absent = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")

    thin_border_side = Side(style="thin", color="D1D5DB")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # ================= ЛИСТ 1: СВОДНЫЙ ТАБЕЛЬ =================
    ws_summary = wb.active
    ws_summary.title = "Сводный табель"
    ws_summary.views.sheetView[0].showGridLines = True

    # Заголовок отчета
    ws_summary.merge_cells("A1:E1")
    ws_summary["A1"] = "ТАБЕЛЬ УЧЕТА ПОСЕЩАЕМОСТИ"
    ws_summary["A1"].font = font_title
    ws_summary["A1"].alignment = align_left

    ws_summary.merge_cells("A2:E2")
    ws_summary["A2"] = f"Период: с {start_date_str} по {end_date_str} (дней: {len(period_dates)})"
    ws_summary["A2"].font = font_subtitle
    ws_summary["A2"].alignment = align_left

    # Формирование шапки таблицы
    headers = ["№", "ФИО Студента"] + period_dates + ["Посещено", "% Посещаемости"]
    header_row_idx = 4

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws_summary.cell(row=header_row_idx, column=col_idx, value=header_text)
        cell.font = font_header
        cell.border = border_cell
        cell.alignment = align_center
        # Выделяем итоговые колонки отдельным цветом
        if col_idx > len(headers) - 2:
            cell.fill = fill_summary_header
        else:
            cell.fill = fill_header

    # Заполнение строк студентов
    sorted_students = sorted(list(all_students))
    for row_offset, student_name in enumerate(sorted_students, start=1):
        current_row = header_row_idx + row_offset
        is_zebra = (row_offset % 2 == 0)
        base_fill = fill_zebra if is_zebra else PatternFill(fill_type=None)

        # №
        c_num = ws_summary.cell(row=current_row, column=1, value=row_offset)
        c_num.font = font_regular
        c_num.alignment = align_center
        c_num.border = border_cell
        if base_fill.fill_type:
            c_num.fill = base_fill

        # ФИО
        c_name = ws_summary.cell(row=current_row, column=2, value=student_name)
        c_name.font = font_bold
        c_name.alignment = align_left
        c_name.border = border_cell
        if base_fill.fill_type:
            c_name.fill = base_fill

        present_days = 0

        # Дни периода
        for d_offset, day_str in enumerate(period_dates, start=1):
            col_target = 2 + d_offset
            record = attendance_map.get((student_name, day_str))
            cell = ws_summary.cell(row=current_row, column=col_target)
            cell.border = border_cell
            cell.alignment = align_center

            if record and record["status"] == "present":
                cell.value = "БЫЛ"
                cell.font = font_present
                cell.fill = fill_present
                present_days += 1
            else:
                cell.value = "Н"
                cell.font = font_absent
                cell.fill = fill_absent

        # Итого посещено
        c_total = ws_summary.cell(row=current_row, column=len(headers) - 1, value=present_days)
        c_total.font = font_bold
        c_total.alignment = align_center
        c_total.border = border_cell

        # Процент посещаемости
        percent_val = round((present_days / len(period_dates)) * 100, 1) if period_dates else 0.0
        c_pct = ws_summary.cell(row=current_row, column=len(headers), value=f"{percent_val}%")
        c_pct.font = font_bold
        c_pct.alignment = align_center
        c_pct.border = border_cell

    # Закрепляем шапку и колонку с ФИО
    ws_summary.freeze_panes = "C5"

    # ================= ЛИСТ 2: ДЕТАЛЬНЫЙ ЖУРНАЛ =================
    ws_detail = wb.create_sheet(title="Детальный журнал")
    ws_detail.views.sheetView[0].showGridLines = True

    detail_headers = ["Дата", "ФИО Студента", "Статус", "Время фиксации"]
    for col_idx, h_text in enumerate(detail_headers, start=1):
        cell = ws_detail.cell(row=1, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.border = border_cell
        cell.alignment = align_center

    for r_idx, (s_name, r_date, r_time, r_status) in enumerate(records, start=2):
        c1 = ws_detail.cell(row=r_idx, column=1, value=r_date)
        c2 = ws_detail.cell(row=r_idx, column=2, value=s_name)
        c3 = ws_detail.cell(row=r_idx, column=3, value="Присутствовал" if r_status == "present" else "Отсутствовал")
        c4 = ws_detail.cell(row=r_idx, column=4, value=r_time)

        for c in (c1, c2, c3, c4):
            c.font = font_regular
            c.border = border_cell
            c.alignment = align_center
        c2.alignment = align_left

        if r_status == "present":
            c3.font = font_present
            c3.fill = fill_present

    # Автоподбор ширины всех колонок для обоих листов
    for sheet in (ws_summary, ws_detail):
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if cell.row < 4 and sheet == ws_summary:
                    continue
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # Увеличенная ширина для колонки ФИО
    ws_summary.column_dimensions["B"].width = 28
    ws_detail.column_dimensions["B"].width = 28

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream